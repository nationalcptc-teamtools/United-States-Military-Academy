"""Offline synthetic reporting starter. Input text never invokes tools or a model."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import tempfile
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from xml.sax.saxutils import escape

VERSION = "0.4.0"
MAX_INPUT_BYTES = 2_000_000
SEVERITIES = ("critical", "high", "medium", "low", "info")
MARKINGS = ("ORCON", "PROPIN", "PII", "PII Red", "FI")
MARKING_POLICY = "cptc.handling/1.0"
CHAPTERS = ("1 Introduction", "2 High-Level Summary and Recommendations",
            "3 Methodology", "4 Detailed Findings", "5 Appendices")
METHOD_SECTIONS = (("information_gathering", "Information gathering"),
                   ("service_enumeration", "Service enumeration"),
                   ("validation", "Validation"), ("cleanup", "Cleanup and retest"))
ROOT = Path(__file__).resolve().parents[2]


class ContractError(ValueError):
    """Error messages name trusted contract locations, never input contents."""


def require(condition, location, rule):
    if not condition:
        raise ContractError(f"{location}: {rule}")


def obj(value, location, required, optional=()):
    require(isinstance(value, dict), location, "expected an object")
    require(set(required) <= value.keys(), location, "required fields missing")
    require(value.keys() <= set(required) | set(optional), location, "unknown fields")


def text(value, location, limit=10000, single_line=False):
    require(isinstance(value, str) and bool(value.strip()), location, "nonempty text required")
    require(len(value) <= limit, location, "text exceeds field limit")
    require(not any(unicodedata.category(c) in ("Cc", "Cf", "Cs") and c != "\n"
                    for c in value), location, "unsupported control characters")
    require(not single_line or "\n" not in value, location, "single-line text required")


def identifier(value, location):
    require(isinstance(value, str) and re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", value),
            location, "expected a stable ID of up to 64 letters, numbers, underscores or hyphens")


def items(value, location, maximum=100, allow_empty=False):
    require(isinstance(value, list) and (allow_empty or len(value) > 0)
            and len(value) <= maximum, location, "invalid list size")


def strings(value, location, limit=2000, maximum=100, single_line=False):
    items(value, location, maximum)
    for i, item in enumerate(value):
        text(item, f"{location}[{i}]", limit, single_line)


def iso_date(value, location):
    require(isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value),
            location, "expected ISO date YYYY-MM-DD")
    try:
        date.fromisoformat(value)
    except ValueError:
        raise ContractError(f"{location}: invalid calendar date") from None


def timestamp(value, location):
    text(value, location, 50, True)
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        require("T" in value and parsed.tzinfo is not None, location,
                "ISO timestamp with timezone required")
    except ValueError:
        raise ContractError(f"{location}: ISO timestamp with timezone required") from None


def artifact_path(value, location):
    text(value, location, 240, True)
    path = PurePosixPath(value)
    require(not path.is_absolute() and "\\" not in value and ":" not in value
            and all(p not in ("", ".", "..") for p in value.split("/")),
            location, "artifact must be a relative path without traversal or URL/drive prefixes")


def validate(data, release=False):
    obj(data, "report", ("schema_version", "data_category", "classification",
        "engagement", "findings"), ("review", "markings"))
    require(data["schema_version"] in ("1.0", "1.1", "1.2"), "report.schema_version", "unsupported schema")
    require(data["data_category"] == "synthetic", "report.data_category",
            "this preparation starter accepts synthetic data only")
    text(data["classification"], "report.classification", 100, True)
    e = data["engagement"]
    structured = data["schema_version"] == "1.2"
    require(structured or "markings" not in data, "report.markings", "requires schema 1.2")
    if structured:
        require("markings" in data, "report.markings", "structured markings required")
        validate_markings(data["markings"], "report.markings")
        require("ORCON" in data["markings"], "report.markings", "ORCON baseline required")
    optional = ("objective", "recommendations", "methodology_sections") if data["schema_version"] != "1.0" else ()
    obj(e, "engagement", ("id", "client", "team", "report_date", "scope", "methodology",
                         "executive_summary", "limitations"), optional)
    identifier(e["id"], "engagement.id")
    for key in ("client", "team"):
        text(e[key], f"engagement.{key}", 120, True)
    iso_date(e["report_date"], "engagement.report_date")
    strings(e["scope"], "engagement.scope")
    text(e["methodology"], "engagement.methodology")
    text(e["executive_summary"], "engagement.executive_summary")
    strings(e["limitations"], "engagement.limitations")
    if "objective" in e:
        text(e["objective"], "engagement.objective")
    if "recommendations" in e:
        strings(e["recommendations"], "engagement.recommendations", maximum=20)
    if "methodology_sections" in e:
        details = e["methodology_sections"]
        obj(details, "engagement.methodology_sections", (), tuple(key for key, _ in METHOD_SECTIONS))
        require(bool(details), "engagement.methodology_sections", "at least one supplied section required")
        for key, value in details.items():
            text(value, f"engagement.methodology_sections.{key}")
    if "review" in data:
        obj(data["review"], "review", ("reviewer", "reviewed_at"))
        text(data["review"]["reviewer"], "review.reviewer", 120, True)
        iso_date(data["review"]["reviewed_at"], "review.reviewed_at")
    if release:
        require("review" in data, "review", "human reviewer and date required for release label")
    items(data["findings"], "findings", maximum=200, allow_empty=True)
    seen_findings, seen_evidence = set(), set()
    for i, f in enumerate(data["findings"]):
        at = f"findings[{i}]"
        obj(f, at, ("id", "title", "severity", "severity_rationale", "status", "assets",
                    "description", "validation_steps", "business_impact", "remediation",
                    "retest_steps", "evidence"), ("cves", "cvss", "markings") if structured else ("cves", "cvss"))
        if "markings" in f:
            validate_markings(f["markings"], f"{at}.markings")
        identifier(f["id"], f"{at}.id")
        require(f["id"] not in seen_findings, f"{at}.id", "duplicate finding ID")
        seen_findings.add(f["id"])
        text(f["title"], f"{at}.title", 180, True)
        require(f["severity"] in SEVERITIES, f"{at}.severity", "invalid severity")
        require(f["status"] in ("candidate", "confirmed"), f"{at}.status", "invalid status")
        require(not release or f["status"] == "confirmed", f"{at}.status",
                "release label requires confirmed findings")
        strings(f["assets"], f"{at}.assets", 300, 100, True)
        for key in ("severity_rationale", "description", "business_impact", "remediation"):
            text(f[key], f"{at}.{key}")
        for key in ("validation_steps", "retest_steps"):
            strings(f[key], f"{at}.{key}", maximum=50)
        items(f["evidence"], f"{at}.evidence", maximum=50)
        for j, evidence in enumerate(f["evidence"]):
            loc = f"{at}.evidence[{j}]"
            obj(evidence, loc, ("id", "observed_at", "collector", "tool", "tool_version", "summary"),
                ("artifact", "sha256", "markings") if structured else ("artifact", "sha256"))
            if "markings" in evidence:
                validate_markings(evidence["markings"], f"{loc}.markings")
            identifier(evidence["id"], f"{loc}.id")
            require(evidence["id"] not in seen_evidence, f"{loc}.id", "duplicate evidence ID")
            seen_evidence.add(evidence["id"])
            timestamp(evidence["observed_at"], f"{loc}.observed_at")
            for key in ("collector", "tool", "tool_version"):
                text(evidence[key], f"{loc}.{key}", 120, True)
            text(evidence["summary"], f"{loc}.summary")
            if "artifact" in evidence:
                artifact_path(evidence["artifact"], f"{loc}.artifact")
            if "sha256" in evidence:
                require("artifact" in evidence, f"{loc}.sha256", "hash requires artifact reference")
                require(isinstance(evidence["sha256"], str)
                        and re.fullmatch(r"[a-f0-9]{64}", evidence["sha256"]),
                        f"{loc}.sha256", "expected lowercase SHA-256")
        if "cves" in f:
            items(f["cves"], f"{at}.cves", maximum=30)
            seen_cves = set()
            for j, cve in enumerate(f["cves"]):
                loc = f"{at}.cves[{j}]"
                obj(cve, loc, ("id", "applicability", "match_basis", "source", "retrieved_at"))
                require(isinstance(cve["id"], str) and re.fullmatch(r"CVE-[0-9]{4}-[0-9]{4,}", cve["id"]),
                        f"{loc}.id", "invalid CVE ID")
                require(cve["id"] not in seen_cves, f"{loc}.id", "duplicate CVE reference")
                seen_cves.add(cve["id"])
                require(cve["applicability"] in ("candidate", "verified"), f"{loc}.applicability",
                        "invalid applicability")
                require(not release or cve["applicability"] == "verified", f"{loc}.applicability",
                        "release label requires verified CVE applicability")
                text(cve["match_basis"], f"{loc}.match_basis", 2000)
                text(cve["source"], f"{loc}.source", 1000, True)
                iso_date(cve["retrieved_at"], f"{loc}.retrieved_at")
        if "cvss" in f:
            c = f["cvss"]
            loc = f"{at}.cvss"
            obj(c, loc, ("version", "score", "vector", "source", "retrieved_at"))
            require(c["version"] in ("3.1", "4.0"), f"{loc}.version", "unsupported CVSS version")
            require(type(c["score"]) in (int, float) and 0 <= c["score"] <= 10
                    and math.isfinite(c["score"]), f"{loc}.score", "finite score from 0 to 10 required")
            text(c["vector"], f"{loc}.vector", 512, True)
            require(c["vector"].startswith(f"CVSS:{c['version']}/")
                    and re.fullmatch(r"CVSS:(3\.1|4\.0)(/[A-Z]{1,4}:[A-Z0-9.]+)+", c["vector"]),
                    f"{loc}.vector", "invalid vector shape or version")
            text(c["source"], f"{loc}.source", 1000, True)
            iso_date(c["retrieved_at"], f"{loc}.retrieved_at")
    return data


def _pairs(pairs):
    result = {}
    for key, value in pairs:
        require(key not in result, "input JSON", "duplicate object key")
        result[key] = value
    return result


def load(raw):
    require(len(raw) <= MAX_INPUT_BYTES, "input", "input exceeds 2 MB limit")
    try:
        return json.loads(raw.decode("utf-8-sig"), object_pairs_hook=_pairs,
                          parse_constant=lambda _: (_ for _ in ()).throw(
                              ContractError("input JSON: non-finite numbers are prohibited")))
    except ContractError:
        raise
    except (ValueError, RecursionError):
        raise ContractError("input JSON: invalid UTF-8 JSON") from None


def md(value):
    # Escape CommonMark punctuation and raw HTML; preserve visible literal text.
    value = str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    punctuation = "\\" + chr(96) + "*_{}[]()#+-.!|~="
    value = "".join("\\" + c if c in punctuation else c for c in value)
    # Entities preserve source spaces without letting indentation create code.
    # Insert them after escaping so their ampersands are not escaped again.
    value = re.sub(r"(?m)^ +", lambda match: "&#32;" * len(match[0]), value)
    return value.replace("\n", "  \n")


def validate_markings(value, location):
    items(value, location, maximum=len(MARKINGS))
    require(all(isinstance(tag, str) and tag in MARKINGS for tag in value),
            location, "unknown handling marking")
    require(len(set(value)) == len(value), location, "duplicate handling marking")


def marking_label(tags):
    tags = set(tags) | {"ORCON"}
    if "PII Red" in tags:
        tags.discard("PII")
    return "//".join(tag for tag in MARKINGS if tag in tags)


def handling(data):
    """Compute restrictions without mutating input or inferring sensitivity from prose."""
    structured = data["schema_version"] == "1.2"
    baseline = {"ORCON"}
    # Known explicit labels can only add restrictions, including during schema migration.
    baseline.update(tag.strip() for tag in data["classification"].split("//")
                    if tag.strip() in MARKINGS)
    if structured:
        baseline.update(data["markings"])
    overall, findings, evidence = set(baseline), {}, {}
    for finding in data["findings"]:
        own = baseline | set(finding.get("markings", []))
        derived = set(own)
        for source in finding["evidence"]:
            source_tags = own | set(source.get("markings", []))
            evidence[source["id"]] = marking_label(source_tags)
            derived.update(source_tags)
        findings[finding["id"]] = marking_label(derived)
        overall.update(derived)
    return {"report": marking_label(overall), "findings": findings, "evidence": evidence,
            "status": "structured" if structured else "legacy_unverified"}


def handling_notice(data):
    if data["schema_version"] != "1.2":
        return ("Legacy input: structured source markings are unavailable; classification inheritance "
                "is not verified. ORCON and recognized supplied legacy tags are retained. "
                "Use schema 1.2 to record source restrictions.")
    return ("ORCON applies to all data and derived analysis. PROPIN marks internal systems, methods "
            "and procedures; PII marks employee/customer information; PII Red marks identity-theft-sensitive "
            "PII; FI marks financial information. Combined analysis retains all applicable restrictions; "
            "PII Red supersedes PII. Markings do not authorize disclosure or determine sensitivity from prose.")


def sections(data):
    """Shared fact sections keep Markdown and PDF aligned."""
    controls = handling(data)
    for f in data["findings"]:
        parts = [
            ("Handling marking", [controls["findings"][f["id"]]]),
            ("Business severity", [f"{f['severity'].upper()} - {f['severity_rationale']}"]),
            ("Validation status", [f["status"].upper()]),
            ("Affected assets", f["assets"]),
            ("Observation", [f["description"]]),
            ("Validation steps and results", f["validation_steps"]),
            ("Business impact", [f["business_impact"]]),
            ("Remediation", [f["remediation"]]),
            ("Retest plan", f["retest_steps"]),
        ]
        for e in f["evidence"]:
            lines = [
                f"Handling marking: {controls['evidence'][e['id']]}",
                f"Observed: {e['observed_at']} | Collector: {e['collector']}",
                f"Source: {e['tool']} {e['tool_version']}",
                f"Normalized evidence summary: {e['summary']}",
            ]
            if "artifact" in e:
                lines.append(f"Private artifact reference (not embedded or verified): {e['artifact']}")
            if "sha256" in e:
                lines.append(f"Supplied artifact SHA-256 (not verified): {e['sha256']}")
            parts.append((f"Evidence {e['id']}", lines))
        if "cves" in f:
            parts.append(("CVE references (supplied)", [
                f"{c['id']} | {c['applicability'].upper()} | Basis: {c['match_basis']} | "
                f"Source: {c['source']} | Retrieved: {c['retrieved_at']}"
                for c in f["cves"]]))
        if "cvss" in f:
            c = f["cvss"]
            parts.append(("CVSS metadata (supplied, not recalculated)", [
                f"CVSS {c['version']} | Score: {c['score']} | Vector: {c['vector']}",
                f"Source: {c['source']} | Retrieved: {c['retrieved_at']}",
                "Published severity metadata does not establish business risk or applicability.",
            ]))
        yield f, parts


def document_blocks(data):
    """One ordered content view for both outputs; no source narrative is inferred."""
    e = data["engagement"]
    yield "heading", (1, CHAPTERS[0])
    yield "heading", (2, "Assessment objective")
    yield "paragraph", e.get("objective", "An assessment objective was not supplied.")
    yield "heading", (2, "Scope")
    yield "list", (e["scope"], False)
    yield "heading", (2, "Limitations")
    yield "list", (e["limitations"], False)
    yield "heading", (2, "Data handling")
    yield "paragraph", f"Report handling marking: {handling(data)['report']}"
    yield "paragraph", handling_notice(data)
    yield "heading", (1, CHAPTERS[1])
    yield "paragraph", e["executive_summary"]
    counts = Counter(f["severity"] for f in data["findings"])
    statuses = Counter(f["status"] for f in data["findings"])
    yield "heading", (2, "Finding summary")
    yield "paragraph", (f"Total: {len(data['findings'])} | Confirmed: {statuses['confirmed']} | "
                        f"Candidate: {statuses['candidate']}")
    yield "paragraph", " | ".join(f"{s.title()}: {counts[s]}" for s in SEVERITIES)
    yield "heading", (2, "Severity method")
    yield "paragraph", ("Supplied business severity; matrix not configured. No likelihood/impact "
                        "scoring is performed. CVSS, when supplied, is separate severity metadata.")
    yield "heading", (2, "Recommendations")
    if "recommendations" in e:
        yield "list", (e["recommendations"], False)
    else:
        yield "paragraph", ("High-level recommendations were not supplied. Refer to the "
                            "individual remediation and retest actions in Section 4.")
    yield "heading", (1, CHAPTERS[2])
    yield "heading", (2, "Assessment approach")
    yield "paragraph", e["methodology"]
    for key, label in METHOD_SECTIONS:
        if key in e.get("methodology_sections", {}):
            yield "heading", (2, label)
            yield "paragraph", e["methodology_sections"][key]
    yield "heading", (1, CHAPTERS[3])
    if not data["findings"]:
        yield "paragraph", ("No findings were supplied. This does not establish that assessed "
                            "assets are secure; review the stated scope and limitations.")
    for finding, parts in sections(data):
        yield "finding", f"{finding['id']}: {finding['title']}"
        for title, values in parts:
            yield "heading", (3, title)
            if title in ("Validation steps and results", "Retest plan"):
                yield "list", (values, True)
            elif title == "Affected assets":
                yield "list", (values, False)
            else:
                for value in values:
                    yield "paragraph", value
    yield "heading", (1, CHAPTERS[4])
    yield "heading", (2, "Evidence register")
    for finding in data["findings"]:
        for evidence in finding["evidence"]:
            yield "paragraph", (f"{evidence['id']} | Finding {finding['id']} | "
                                f"{evidence['observed_at']} | Collector: {evidence['collector']} | "
                                f"Source: {evidence['tool']} {evidence['tool_version']}")
    if not data["findings"]:
        yield "paragraph", "No evidence records were supplied."
    yield "heading", (2, "Review record")
    if "review" in data:
        yield "paragraph", (f"Supplied reviewer: {data['review']['reviewer']} | "
                            f"Supplied review date: {data['review']['reviewed_at']}")
    else:
        yield "paragraph", "Human review pending."
    yield "paragraph", "Review labels are local metadata. Inspect the actual artifacts before any authorized use."
    yield "paragraph", ("Evidence files are referenced, never opened. Supplied hashes do not verify "
                        "the observations; the manifest identifies the input and output files.")
    yield "paragraph", f"Generator {VERSION} | Input schema {data['schema_version']}"


def render_markdown(data, release=False):
    e = data["engagement"]
    label = "REVIEWED TRAINING COPY" if release else "DRAFT"
    lines = ["# Penetration Testing Report", "", f"**{label} | SYNTHETIC DATA**", "",
             f"Client: {md(e['client'])}  \nTeam: {md(e['team'])}  \n"
             f"Engagement: {md(e['id'])}  \nReport date: {md(e['report_date'])}  \n"
             f"Supplied document label: {md(data['classification'])}  \n"
             f"Handling marking: {handling(data)['report']}", "",
             "Training demonstration only. No real assessment, event permission, or submission is established.",
             "", "## Contents", ""]
    for chapter in CHAPTERS:
        anchor = chapter.lower().replace(" ", "-")
        lines.append(f"- [{chapter}](#{anchor})")
    lines.append("")
    for kind, value in document_blocks(data):
        if kind == "heading":
            level, title = value
            lines += ["#" * (level + 1) + " " + md(title), ""]
        elif kind == "finding":
            lines += ["### " + md(value), ""]
        elif kind == "list":
            values, numbered = value
            for index, item in enumerate(values, 1):
                prefix = f"{index}. " if numbered else "- "
                # Keep all physical lines inside this item without interpreting source indentation.
                lines.append(prefix + md(item).replace("\n", "\n" + " " * len(prefix)))
            lines.append("")
        else:
            lines += [md(value), ""]
    return "\n".join(lines)


def all_strings(value):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for item in value.values():
            yield from all_strings(item)
    elif isinstance(value, list):
        for item in value:
            yield from all_strings(item)


def render_pdf(data, destination, release=False):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
        from reportlab.platypus.tableofcontents import TableOfContents
    except ImportError:
        raise ContractError("PDF: install the declared ReportLab dependency in an approved environment") from None

    candidates = [
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
        (Path("/Library/Fonts/Arial.ttf"), Path("/Library/Fonts/Arial Bold.ttf")),
    ]
    normal, bold = "Helvetica", "Helvetica-Bold"
    glyphs = set(range(32, 256))
    for regular_path, bold_path in candidates:
        if regular_path.is_file() and bold_path.is_file():
            normal, bold = "ReportText", "ReportBold"
            pdfmetrics.registerFont(TTFont(normal, str(regular_path)))
            pdfmetrics.registerFont(TTFont(bold, str(bold_path)))
            glyphs = set(pdfmetrics.getFont(normal).face.charToGlyph) & set(
                pdfmetrics.getFont(bold).face.charToGlyph)
            break
    require(all(ord(char) in glyphs or char == "\n" for s in all_strings(data) for char in s),
            "PDF", "input contains glyphs unavailable in the installed fonts; use Markdown or approved fonts")
    navy, ink, muted = colors.HexColor("#17365D"), colors.HexColor("#222222"), colors.HexColor("#526575")
    body = ParagraphStyle("Body", fontName=normal, fontSize=10.5, leading=14.5, textColor=ink,
                          spaceAfter=6, alignment=TA_LEFT, splitLongWords=True)
    styles = {
        "body": body,
        "title": ParagraphStyle("Title", parent=body, fontName=bold, fontSize=26, leading=32,
                                textColor=navy, spaceAfter=18),
        "h1": ParagraphStyle("H1", parent=body, fontName=bold, fontSize=17, leading=22,
                             textColor=navy, spaceAfter=12, keepWithNext=True),
        "h2": ParagraphStyle("H2", parent=body, fontName=bold, fontSize=13, leading=18,
                             textColor=navy, spaceBefore=10, spaceAfter=6, keepWithNext=True),
        "h3": ParagraphStyle("H3", parent=body, fontName=bold, fontSize=10.5, leading=14.5,
                             textColor=ink, spaceBefore=5, spaceAfter=3, keepWithNext=True),
        "small": ParagraphStyle("Small", parent=body, fontSize=8.5, leading=12, textColor=muted),
    }

    def p(value, style="body"):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), styles[style])

    def prose(value):
        # Bounded chunks allow a heading and the start of a long narrative to
        # share a page, without keeping the entire narrative with the heading.
        value = str(value)
        chunks = []
        while len(value) > 1600:
            cut = value.rfind(" ", 800, 1600)
            if cut < 0:
                cut = 1600
            chunks.append(p(value[:cut]))
            value = value[cut:]
        chunks.append(p(value))
        return chunks

    e = data["engagement"]
    label = "REVIEWED TRAINING COPY" if release else "DRAFT"
    report_controls = handling(data)
    story = [Spacer(1, 42), p(f"{label} | SYNTHETIC DATA", "small"), Spacer(1, 20),
             p("Penetration Testing Report", "title"), p(e["client"], "h1"), Spacer(1, 18)]
    metadata = [
        ["Engagement", e["id"]],
        ["Team", e["team"]],
        ["Report date", e["report_date"]],
        ["Document label", data["classification"]],
        ["Handling marking", report_controls["report"]],
    ]
    table = Table([[p(key, "small"), p(value)] for key, value in metadata], colWidths=[108, 360])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EDF2F6")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story += [table, Spacer(1, 18),
              p("Training demonstration only. No real assessment, event permission, or submission is established.",
                "small"), PageBreak(), p("Contents", "title")]
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("ContentsEntry", parent=body, fontSize=11,
                                      leading=20, rightIndent=20, spaceBefore=5)]
    toc.dotsMinLevel = 0
    story.append(toc)
    first_finding = True
    for kind, value in document_blocks(data):
        if kind == "heading":
            level, title = value
            if level == 1:
                story.append(PageBreak())
            heading = p(title, f"h{level}")
            if level == 1:
                # Chapter labels are trusted constants. No source markup enters TOC parsing.
                heading._toc_entry = title
            story.append(heading)
        elif kind == "finding":
            if not first_finding:
                story.append(PageBreak())
            first_finding = False
            story.append(p(value, "h2"))
        elif kind == "list":
            values, numbered = value
            for index, item in enumerate(values, 1):
                prefix = f"{index}. " if numbered else "- "
                story += prose(prefix + item)
        else:
            story += prose(value)

    def page_frame(c, doc):
        c.saveState()
        c.setStrokeColor(colors.HexColor("#B6C5D1"))
        c.line(72, 741, 540, 741)
        c.setFont(bold, 8)
        c.setFillColor(navy)
        c.drawString(72, 754, f"{label} | SYNTHETIC")
        c.setFont(normal, 8)
        c.drawRightString(540, 754, "CPTC reporting")
        c.line(72, 54, 540, 54)
        c.setFillColor(muted)
        c.drawString(72, 38, "Synthetic training artifact | Human review required")
        c.drawRightString(540, 38, f"Page {doc.page}")
        c.setFont(bold, 8)
        legacy = " | LEGACY MARKINGS UNVERIFIED" if data["schema_version"] != "1.2" else ""
        c.drawString(72, 24, report_controls["report"] + legacy)
        c.restoreState()

    class StableCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            kwargs["invariant"] = 1
            super().__init__(*args, **kwargs)

    class ReportingDocTemplate(SimpleDocTemplate):
        def afterFlowable(self, flowable):
            if hasattr(flowable, "_toc_entry"):
                self.notify("TOCEntry", (0, flowable._toc_entry, self.page))

    document = ReportingDocTemplate(str(destination), pagesize=letter, rightMargin=72, leftMargin=72,
                                 topMargin=72, bottomMargin=72, title="Synthetic penetration testing report",
                                 author="CPTC reporting starter", pageCompression=1)
    document.multiBuild(story, onFirstPage=page_frame, onLaterPages=page_frame,
                        canvasmaker=StableCanvas, maxPasses=5)


def render_docx(data, destination, release=False):
    """Editable Word export; source strings are never interpreted as OOXML or fields."""
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        from lxml import etree
    except ImportError:
        raise ContractError("Word: install the declared python-docx dependency in an approved environment") from None
    from datetime import timezone
    from io import BytesIO
    from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

    require(all(ord(char) not in (0xFFFE, 0xFFFF) for value in all_strings(data) for char in value),
            "Word", "input contains characters unavailable in XML documents")
    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    doc.settings.odd_and_even_pages_header_footer = False
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.35)
    for name, size in (("Normal", 10.5), ("Title", 26), ("Subtitle", 16),
                       ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 10.5),
                       ("Header", 8), ("Footer", 8), ("List Paragraph", 10.5)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(8)
    doc.styles["Normal"].paragraph_format.widow_control = True
    controls = handling(data)
    label = "REVIEWED TRAINING COPY" if release else "DRAFT"
    header = section.header.paragraphs[0]
    header.add_run(f"{label} | SYNTHETIC DATA").bold = True
    header.add_run("\n" + controls["report"]).bold = True
    if controls["status"] == "legacy_unverified":
        header.add_run(" | LEGACY MARKINGS UNVERIFIED")
    footer = section.footer.paragraphs[0]
    footer.add_run("Synthetic training artifact | Human review required | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def paragraph(value, style=None):
        result = doc.add_paragraph(style=style)
        result.add_run(str(value))
        result.paragraph_format.keep_together = False
        return result

    e = data["engagement"]
    paragraph("Penetration Testing Report", "Title")
    paragraph(e["client"], "Subtitle")
    for key, value in (("Engagement", e["id"]), ("Team", e["team"]),
                       ("Report date", e["report_date"]), ("Document label", data["classification"]),
                       ("Handling marking", controls["report"])):
        row = doc.add_paragraph()
        row.add_run(key + ": ").bold = True
        row.add_run(value)
    paragraph("Training demonstration only. No real assessment, event permission, or submission is established.")
    doc.add_page_break()
    paragraph("Contents", "Heading 1")
    for index, chapter in enumerate(CHAPTERS, 1):
        row = doc.add_paragraph()
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), f"chapter_{index}")
        run, properties, value = OxmlElement("w:r"), OxmlElement("w:rPr"), OxmlElement("w:t")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.append(underline)
        value.text = chapter
        run.extend((properties, value))
        link.append(run)
        row._p.append(link)

    numbering = doc.part.numbering_part.element
    next_abstract = 1 + max(int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum")))
    next_number = 1 + max(int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num")))

    def list_number(numbered):
        nonlocal next_abstract, next_number
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(next_abstract))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("start", "1"), ("numFmt", "decimal" if numbered else "bullet"),
                           ("suff", "space"), ("lvlText", "%1." if numbered else "-")):
            node = OxmlElement("w:" + tag)
            node.set(qn("w:val"), value)
            level.append(node)
        abstract.append(level)
        position = next((i for i, node in enumerate(numbering) if node.tag == qn("w:num")), len(numbering))
        numbering.insert(position, abstract)
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(next_number))
        reference = OxmlElement("w:abstractNumId")
        reference.set(qn("w:val"), str(next_abstract))
        number.append(reference)
        numbering.append(number)
        result = next_number
        next_abstract += 1
        next_number += 1
        return result

    first_finding = True
    for kind, value in document_blocks(data):
        if kind == "heading":
            level, title = value
            row = paragraph(title, f"Heading {level}")
            if level == 1:
                row.paragraph_format.page_break_before = True
                index = CHAPTERS.index(title) + 1
                start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
                start.set(qn("w:id"), str(index))
                start.set(qn("w:name"), f"chapter_{index}")
                end.set(qn("w:id"), str(index))
                row._p.insert(1, start)
                row._p.append(end)
        elif kind == "finding":
            row = paragraph(value, "Heading 2")
            row.paragraph_format.page_break_before = not first_finding
            first_finding = False
        elif kind == "list":
            values, numbered = value
            number_id = list_number(numbered)
            for item in values:
                row = paragraph(item, "List Paragraph")
                row.paragraph_format.left_indent = Inches(0.25)
                row.paragraph_format.first_line_indent = Inches(-0.15)
                properties = row._p.get_or_add_pPr().get_or_add_numPr()
                properties.get_or_add_ilvl().val = 0
                properties.get_or_add_numId().val = number_id
        else:
            paragraph(value)
    properties = doc.core_properties
    properties.title, properties.author = "Synthetic penetration testing report", "CPTC reporting starter"
    properties.last_modified_by = properties.comments = properties.keywords = ""
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    properties.revision = 1
    original, canonical = BytesIO(), BytesIO()
    doc.save(original)
    with ZipFile(original) as source, ZipFile(canonical, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            if name.startswith("customXml/") or name == "docProps/thumbnail.jpeg":
                continue
            raw = source.read(name)
            if name.endswith((".xml", ".rels")):
                root = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True))
                for node in root.iter():
                    for key in list(node.attrib):
                        if etree.QName(key).localname.startswith("rsid"):
                            del node.attrib[key]
                    for child in list(node):
                        relation = child.get("Type", "")
                        part = child.get("PartName", "")
                        if relation.endswith(("/customXml", "/metadata/thumbnail")) or part.startswith("/customXml/"):
                            node.remove(child)
                raw = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type, info.create_system = ZIP_DEFLATED, 0
            target.writestr(info, raw)
    Path(destination).write_bytes(canonical.getvalue())


def sha256(raw):
    return hashlib.sha256(raw).hexdigest()


def generate(raw, output_dir, pdf=False, release=False, *, docx=False):
    data = validate(load(raw), release)
    output_dir = Path(output_dir).resolve()
    require(not any(output_dir == blocked or blocked in output_dir.parents
                    for blocked in (ROOT / "sources", ROOT / "tooling")),
            "destination", "outputs cannot be written under sources or reusable tooling")
    require(not output_dir.exists(), "destination", "already exists; choose a new output directory")
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".report-staging-", dir=output_dir.parent) as staging:
        stage = Path(staging)
        (stage / "report.md").write_text(render_markdown(data, release), encoding="utf-8")
        if pdf:
            render_pdf(data, stage / "report.pdf", release)
        if docx:
            render_docx(data, stage / "report.docx", release)
        artifact_hashes = {p.name: sha256(p.read_bytes()) for p in sorted(stage.iterdir())}
        manifest = {
            "generator_version": VERSION, "schema_version": data["schema_version"],
            "data_category": data["data_category"], "classification": data["classification"],
            "review_label": "reviewed_training_copy" if release else "draft",
            "input_sha256": sha256(raw), "artifacts": artifact_hashes,
            "evidence_files_opened": False, "cvss_recalculated": False,
            "marking_policy": MARKING_POLICY, "effective_markings": handling(data)["report"],
            "classification_inheritance": handling(data)["status"],
        }
        (stage / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
        # Refuse a destination that appeared while rendering; never intentionally overwrite it.
        require(not output_dir.exists(), "destination", "appeared during generation; choose a new directory")
        stage.rename(output_dir)
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="UTF-8 normalized JSON file, or - for stdin")
    parser.add_argument("--output-dir", required=True, help="new private output directory; never overwritten")
    parser.add_argument("--pdf", action="store_true", help="also render PDF using locally installed ReportLab")
    parser.add_argument("--docx", action="store_true", help="also render editable Word using locally installed python-docx")
    parser.add_argument("--release", action="store_true", help="reviewed training label; enforce review metadata")
    args = parser.parse_args(argv)
    try:
        if args.input == "-":
            raw = sys.stdin.buffer.read(MAX_INPUT_BYTES + 1)
        else:
            with Path(args.input).open("rb") as stream:
                raw = stream.read(MAX_INPUT_BYTES + 1)
        manifest = generate(raw, args.output_dir, args.pdf, args.release, docx=args.docx)
    except ContractError as exc:
        print(f"Report rejected: {exc}", file=sys.stderr)
        return 2
    except Exception:
        # Do not expose paths or document text from dependency/OS exceptions.
        print("Report failed: could not read input or render/publish artifacts; check local dependencies "
              "and filesystem permissions. No report destination was intentionally overwritten.", file=sys.stderr)
        return 1
    formats = "Markdown" + (" + PDF" if args.pdf else "") + (" + Word" if args.docx else "")
    print(f"Generated {formats}; "
          f"label={manifest['review_label']}; input_sha256={manifest['input_sha256']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
